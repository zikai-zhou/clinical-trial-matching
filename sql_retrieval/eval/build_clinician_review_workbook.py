#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
build_clinician_review_workbook.py

Build a clinician-facing Excel workbook from the sampled UNION rows produced by
sample_physician_validation_union.py.

Workflow
--------
  1) sample_physician_validation_union.py decides the sampled review set
  2) this script reads sampled_items.csv as the authoritative review list
  3) this script enriches each sampled row with:
       - patient note
       - full structured trial listing
       - retrieval objective (clinician-facing)
       - relevance definition / instructions
       - parsed evidence from shared pair cache
       - clinician annotation columns

Key design choice
-----------------
This script does NOT re-expand union rows into separate SMT / TrialGPT rows.
One sampled CSV row becomes exactly one workbook review row.

Expected sampled CSV
--------------------
Produced by sample_physician_validation_union.py in UNION mode, e.g.:
  sample_id
  mode
  bucket
  patient_id
  trial_id
  union_bucket_used_for_sampling
  ...

Expected shared pair cache layout
---------------------------------
  <shared_pair_cache_root>/<mode>/<patient_id>/<trial_id>/
    relevance.txt
    eligibility.txt

Workbook sheets
---------------
Review sheets:
- chief_complaint_treating_review
- any_condition_treating_review
- any_condition_relevant_review

Reference sheets:
- README
- column_guide
- decision_guide
- retrieval_objectives
- prompts_ccr
- prompts_all
- prompts_all_explore
"""

from __future__ import annotations

import argparse
import ast
import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


MODES = ["ccr", "all", "all-explore"]

MODE_TO_SHEET = {
    "ccr": "chief_complaint_treating_review",
    "all": "any_condition_treating_review",
    "all-explore": "any_condition_relevant_review",
}

MODE_TO_DISPLAY = {
    "ccr": "Chief-complaint-treating",
    "all": "Any-condition-treating",
    "all-explore": "Any-condition-relevant",
}

MODE_TO_DESCRIPTION = {
    "ccr": (
        "Trials intended to directly treat the patient's chief complaint. "
        "This is the narrowest retrieval objective."
    ),
    "all": (
        "Trials intended to treat any clinically important condition described "
        "in the patient case, not only the chief complaint."
    ),
    "all-explore": (
        "Trials that are clinically relevant to any condition described in the "
        "patient case, including non-treatment relationships when appropriate. "
        "This is the broadest retrieval objective."
    ),
}

CLINICIAN_DECISION_CHOICES = [
    "accept",
    "reject_relevance",
    "reject_eligibility",
    "reject_both",
    "uncertain",
]


# =============================================================================
# Basic IO
# =============================================================================

def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def load_jsonl_as_dict(path: Path, id_key: str = "_id") -> Dict[str, dict]:
    out: Dict[str, dict] = {}
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            obj = json.loads(line)
            k = obj.get(id_key)
            if not isinstance(k, str):
                raise ValueError(f"Missing id_key={id_key} in {path}")
            out[k] = obj
    return out


def read_csv_rows(path: Path) -> List[dict]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def normalize_str(x: Any) -> str:
    if x is None:
        return ""
    return str(x)


def pretty_json(obj: Any) -> str:
    try:
        return json.dumps(obj, ensure_ascii=False, indent=2)
    except Exception:
        return str(obj)


# =============================================================================
# Parsing helpers
# =============================================================================

def _coerce_boolish_str(s: str) -> Optional[bool]:
    sl = s.strip().lower()
    if sl in ("true", "yes", "y", "relevant", "eligible"):
        return True
    if sl in ("false", "no", "n", "irrelevant", "not relevant", "ineligible", "not eligible"):
        return False
    return None


def _extract_tagged_payload(text: str, tag: str) -> Optional[str]:
    s = (text or "").strip()
    m = re.search(
        rf"<\s*{re.escape(tag)}\s*>\s*(.*?)\s*<\s*/\s*{re.escape(tag)}\s*(?:>|$)",
        s,
        flags=re.DOTALL | re.IGNORECASE,
    )
    if not m:
        return None
    return m.group(1).strip()


def _json_parse_best_effort(payload: str) -> Any:
    try:
        return json.loads(payload)
    except Exception:
        return ast.literal_eval(payload)


def _json_leading_value(text: str) -> Tuple[Optional[Any], Optional[int]]:
    s = (text or "").lstrip()
    if not s:
        return None, None
    dec = json.JSONDecoder()
    try:
        obj, idx = dec.raw_decode(s)
        return obj, idx
    except Exception:
        return None, None


def _json_value_after_marker(text: str, marker_regex: str) -> Optional[Any]:
    m = re.search(marker_regex, text or "", flags=re.IGNORECASE)
    if not m:
        return None
    tail = (text or "")[m.end():].lstrip()
    if not tail:
        return None
    dec = json.JSONDecoder()
    try:
        obj, _idx = dec.raw_decode(tail)
        return obj
    except Exception:
        return None


def _json_last_list_anywhere(text: str) -> Optional[list]:
    positions = [m.start() for m in re.finditer(r"\[", text or "")]
    dec = json.JSONDecoder()
    for pos in reversed(positions[-200:]):
        cand = (text or "")[pos:].strip()
        if not cand:
            continue
        if cand.startswith("[]"):
            return []
        try:
            obj, _idx = dec.raw_decode(cand)
        except Exception:
            continue
        if isinstance(obj, list):
            return obj
    return None


def _extract_bool_from_obj(obj: Any, keys: List[str]) -> Optional[bool]:
    if isinstance(obj, bool):
        return obj
    if isinstance(obj, str):
        return _coerce_boolish_str(obj)

    if isinstance(obj, dict):
        for k in keys:
            if k in obj:
                v = obj[k]
                if isinstance(v, bool):
                    return v
                if isinstance(v, str):
                    b = _coerce_boolish_str(v)
                    if b is not None:
                        return b
                if isinstance(v, list):
                    return len(v) > 0
        for v in obj.values():
            b = _extract_bool_from_obj(v, keys)
            if b is not None:
                return b
        return None

    return None


def _list_any_true_semantics(
    obj_list: list,
    *,
    decision_key_candidates: List[str],
    true_tokens: List[str],
) -> Optional[bool]:
    if not isinstance(obj_list, list):
        return None
    if len(obj_list) == 0:
        return False

    if all(isinstance(x, str) for x in obj_list):
        return True

    if all(isinstance(x, dict) for x in obj_list):
        saw_any = False
        saw_true = False
        for d in obj_list:
            for k in decision_key_candidates:
                if k not in d:
                    continue
                v = d[k]
                saw_any = True
                if isinstance(v, bool):
                    if v:
                        saw_true = True
                elif isinstance(v, str):
                    if v.strip().lower() in true_tokens:
                        saw_true = True
        if saw_any:
            return True if saw_true else False
        return None

    return None


def parse_relevance_bool(text: str) -> Optional[bool]:
    payload = _extract_tagged_payload(text, "relevant_subcohorts")
    if payload is not None:
        try:
            obj = _json_parse_best_effort(payload)
            if isinstance(obj, list):
                if len(obj) == 0:
                    return False
                for entry in obj:
                    if isinstance(entry, dict):
                        name = entry.get("subcohort_name")
                        if isinstance(name, str) and name.strip():
                            return True
                for entry in obj:
                    if isinstance(entry, dict):
                        name2 = entry.get("subcohortName") or entry.get("name")
                        if isinstance(name2, str) and name2.strip():
                            return True
                b = _list_any_true_semantics(
                    obj,
                    decision_key_candidates=["relevant", "is_relevant", "decision", "relevance_decision"],
                    true_tokens=["relevant", "true", "yes", "y"],
                )
                if b is not None:
                    return b
        except Exception:
            pass

    obj0, _ = _json_leading_value(text)
    if isinstance(obj0, list):
        b = _list_any_true_semantics(
            obj0,
            decision_key_candidates=["relevant", "is_relevant", "decision", "relevance_decision"],
            true_tokens=["relevant", "true", "yes", "y"],
        )
        if b is not None:
            return b

    obj1 = _json_value_after_marker(text, r"\boutput\s*:\s*")
    if isinstance(obj1, list):
        b = _list_any_true_semantics(
            obj1,
            decision_key_candidates=["relevant", "is_relevant", "decision", "relevance_decision"],
            true_tokens=["relevant", "true", "yes", "y"],
        )
        if b is not None:
            return b

    obj2 = _json_last_list_anywhere(text)
    if isinstance(obj2, list):
        b = _list_any_true_semantics(
            obj2,
            decision_key_candidates=["relevant", "is_relevant", "decision", "relevance_decision"],
            true_tokens=["relevant", "true", "yes", "y"],
        )
        if b is not None:
            return b

    try:
        obj = json.loads((text or "").strip())
        b = _extract_bool_from_obj(obj, ["relevant", "is_relevant", "relevance", "decision", "relevance_decision"])
        if b is not None:
            return b
    except Exception:
        pass

    tl = (text or "").lower()
    if "not relevant" in tl or "irrelevant" in tl:
        return False
    if re.search(r"(^|\W)relevant(\W|$)", tl) and not re.search(r"\bnot\s+relevant\b", tl):
        return True
    return None


def parse_eligibility_bool(text: str) -> Optional[bool]:
    s = (text or "").strip()

    payload = _extract_tagged_payload(s, "subcohort_eligibility_decisions")
    if payload is not None:
        try:
            obj = _json_parse_best_effort(payload)
            if isinstance(obj, list):
                b = _list_any_true_semantics(
                    obj,
                    decision_key_candidates=[
                        "eligibility_decision",
                        "eligibilityDecision",
                        "decision",
                        "eligible",
                        "is_eligible",
                    ],
                    true_tokens=["eligible", "true", "yes", "y"],
                )
                if b is not None:
                    return b
        except Exception:
            pass

    payload = _extract_tagged_payload(s, "eligible_subcohorts")
    if payload is not None:
        try:
            obj = _json_parse_best_effort(payload)
            if isinstance(obj, list):
                b = _list_any_true_semantics(
                    obj,
                    decision_key_candidates=["eligible", "is_eligible", "decision", "eligibility_decision"],
                    true_tokens=["eligible", "true", "yes", "y"],
                )
                if b is not None:
                    return b
        except Exception:
            pass

    obj0, _ = _json_leading_value(s)
    if isinstance(obj0, list):
        b = _list_any_true_semantics(
            obj0,
            decision_key_candidates=["eligibility_decision", "decision", "eligible", "is_eligible"],
            true_tokens=["eligible", "true", "yes", "y"],
        )
        if b is not None:
            return b

    obj1 = _json_value_after_marker(s, r"\boutput\s*:\s*")
    if isinstance(obj1, list):
        b = _list_any_true_semantics(
            obj1,
            decision_key_candidates=["eligibility_decision", "decision", "eligible", "is_eligible"],
            true_tokens=["eligible", "true", "yes", "y"],
        )
        if b is not None:
            return b

    obj2 = _json_last_list_anywhere(s)
    if isinstance(obj2, list):
        b = _list_any_true_semantics(
            obj2,
            decision_key_candidates=["eligibility_decision", "decision", "eligible", "is_eligible"],
            true_tokens=["eligible", "true", "yes", "y"],
        )
        if b is not None:
            return b

    try:
        obj = json.loads(s)
        b = _extract_bool_from_obj(
            obj,
            ["eligible", "is_eligible", "eligibility", "decision", "eligibility_decision", "eligibilityDecision"],
        )
        if b is not None:
            return b
    except Exception:
        pass

    tl = s.lower()
    if re.search(r"\bno\s+relevant\s+subcohorts?\b", tl):
        return False
    if re.search(r"\bno\s+eligibility\s+decisions?\s+are\s+required\b", tl):
        return False
    if re.search(r"\bno\s+eligibility\s+assessment\s+is\s+required\b", tl):
        return False
    if re.search(r"\bthere\s+are\s+no\s+subcohorts?\s+to\s+evaluate\b", tl):
        return False

    if "ineligible" in tl or "not eligible" in tl:
        return False
    if re.search(r"(^|\W)eligible(\W|$)", tl) and "ineligible" not in tl and "not eligible" not in tl:
        return True
    return None


def parse_relevant_subcohorts_obj(text: str) -> Any:
    payload = _extract_tagged_payload(text, "relevant_subcohorts")
    if payload is None:
        return None
    try:
        return _json_parse_best_effort(payload)
    except Exception:
        return payload


def parse_eligibility_obj(text: str) -> Any:
    payload = _extract_tagged_payload(text, "subcohort_eligibility_decisions")
    if payload is not None:
        try:
            return _json_parse_best_effort(payload)
        except Exception:
            return payload

    payload2 = _extract_tagged_payload(text, "eligible_subcohorts")
    if payload2 is not None:
        try:
            return _json_parse_best_effort(payload2)
        except Exception:
            return payload2

    return None


# =============================================================================
# Prompt loading
# =============================================================================

def load_prompt_components(prompt_dir: Path) -> Dict[str, Dict[str, str]]:
    relevance_base = read_text(prompt_dir / "relevance_base.prompt")
    eligibility_template = read_text(prompt_dir / "eligibility.prompt")

    out: Dict[str, Dict[str, str]] = {}
    for mode in MODES:
        def_path = prompt_dir / f"relevance_def_{mode}.prompt"
        instr_path = prompt_dir / f"relevance_instructions_{mode}.prompt"
        out[mode] = {
            "relevance_base": relevance_base,
            "relevance_def": read_text(def_path),
            "relevance_instructions": read_text(instr_path) if instr_path.exists() else "",
            "eligibility_template": eligibility_template,
        }
    return out


# =============================================================================
# Review items
# =============================================================================

@dataclass(frozen=True)
class ReviewItem:
    sample_id: str
    mode: str
    patient_id: str
    trial_id: str
    union_bucket_used_for_sampling: str


def build_review_items(sample_rows: List[dict]) -> List[ReviewItem]:
    out: List[ReviewItem] = []
    seen = set()

    for row in sample_rows:
        sample_id = normalize_str(row.get("sample_id")).strip()
        mode = normalize_str(row.get("mode")).strip()
        patient_id = normalize_str(row.get("patient_id")).strip()
        trial_id = normalize_str(row.get("trial_id") or row.get("parent_trial_id")).strip()

        if not mode or not patient_id or not trial_id:
            continue

        key = (mode, patient_id, trial_id)
        if key in seen:
            continue
        seen.add(key)

        union_bucket_used_for_sampling = normalize_str(
            row.get("union_bucket_used_for_sampling") or row.get("bucket") or row.get("llm_label_bucket")
        ).strip()

        out.append(
            ReviewItem(
                sample_id=sample_id,
                mode=mode,
                patient_id=patient_id,
                trial_id=trial_id,
                union_bucket_used_for_sampling=union_bucket_used_for_sampling,
            )
        )

    return out


# =============================================================================
# Shared pair cache helpers
# =============================================================================

def get_pair_cache_trial_dir(
    *,
    shared_pair_cache_root: Path,
    mode: str,
    patient_id: str,
    trial_id: str,
) -> Path:
    return shared_pair_cache_root / mode / patient_id / trial_id


# =============================================================================
# Trial listing formatting
# =============================================================================

def build_trial_listing(trial_obj: dict) -> str:
    if not trial_obj:
        return "N/A"

    metadata = trial_obj.get("metadata") or {}

    def s(x: Any) -> str:
        if x is None:
            return "N/A"
        if isinstance(x, list):
            return ", ".join(str(v) for v in x) if x else "N/A"
        sx = str(x).strip()
        return sx if sx else "N/A"

    parts = [
        f"trial_id: {s(trial_obj.get('_id'))}",
        f"title: {s(trial_obj.get('title'))}",
        f"brief_title: {s(metadata.get('brief_title'))}",
        f"phase: {s(metadata.get('phase'))}",
        f"enrollment: {s(metadata.get('enrollment'))}",
        f"diseases: {s(metadata.get('diseases_list') or metadata.get('diseases'))}",
        f"drugs: {s(metadata.get('drugs_list') or metadata.get('drugs'))}",
        "",
        "brief_summary:",
        s(metadata.get("brief_summary")),
        "",
        "full_text:",
        s(trial_obj.get("text")),
        "",
        "inclusion_criteria:",
        s(metadata.get("inclusion_criteria")),
        "",
        "exclusion_criteria:",
        s(metadata.get("exclusion_criteria")),
        "",
        "metadata_full_json:",
        pretty_json(metadata) if metadata else "N/A",
    ]
    return "\n".join(parts).strip()


# =============================================================================
# Pretty aggregation
# =============================================================================

def extract_optional_relevance_rationale(
    relevant_subcohorts_obj: Any,
    relevance_output_raw: str,
) -> str:
    # Preferred source: explicit relevance summary block in raw relevance output
    summary_payload = _extract_tagged_payload(relevance_output_raw, "relevance_summary")
    if summary_payload and summary_payload.strip():
        return summary_payload.strip()

    # Backward-compatible fallback: try older structured fields if present
    obj = relevant_subcohorts_obj

    if isinstance(obj, dict):
        for k in ("relevance_summary", "relevance_rationale"):
            v = obj.get(k)
            if isinstance(v, str) and v.strip():
                return v.strip()

    if isinstance(obj, list):
        reasons = []
        for entry in obj:
            if isinstance(entry, dict):
                for k in ("relevance_summary", "relevance_rationale", "why_relevant"):
                    v = entry.get(k)
                    if isinstance(v, str) and v.strip():
                        reasons.append(v.strip())
                        break
        if reasons:
            return "\n---\n".join(reasons)

    return ""

# =============================================================================
# Workbook styling
# =============================================================================

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)

LLM_OUTPUT_FILL = PatternFill("solid", fgColor="EAF2F8")
LLM_OUTPUT_FONT = Font(color="000000", bold=True)

CLINICIAN_INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
CLINICIAN_INPUT_FONT = Font(color="000000", bold=True)

CLINICIAN_INPUT_BODY_FILL = PatternFill("solid", fgColor="FFF8E1")

TOP_WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(vertical="center", horizontal="center")


def style_header(ws, row_idx: int = 1) -> None:
    for cell in ws[row_idx]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER


def style_semantic_columns(ws) -> None:
    header_to_col = {
        normalize_str(cell.value): idx
        for idx, cell in enumerate(ws[1], start=1)
    }

    llm_output_headers = {
        "any_trial_relevant_from_llm_judge",
        "relevant_subcohorts",
        "any_trial_relevant_and_eligible_from_llm_judge",
        "relevant_subcohorts_eligibility_determination",
        "relevance_rationale",
    }

    clinician_input_headers = {
        "clinician_decision",
        "clinician_rationale",
    }

    for header, col_idx in header_to_col.items():
        cell = ws.cell(row=1, column=col_idx)
        if header in llm_output_headers:
            cell.fill = LLM_OUTPUT_FILL
            cell.font = LLM_OUTPUT_FONT
            cell.alignment = CENTER
        elif header in clinician_input_headers:
            cell.fill = CLINICIAN_INPUT_FILL
            cell.font = CLINICIAN_INPUT_FONT
            cell.alignment = CENTER
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row=row_idx, column=col_idx).fill = CLINICIAN_INPUT_BODY_FILL


def auto_width_with_caps(ws, caps: Optional[Dict[str, int]] = None) -> None:
    caps = caps or {}
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = 0
        header = normalize_str(ws.cell(row=1, column=col_idx).value)
        for cell in col_cells:
            val = "" if cell.value is None else str(cell.value)
            effective_len = 0 if val == "N/A" else len(val)
            max_len = max(max_len, effective_len)
            cell.alignment = TOP_WRAP
        width = min(max_len + 2, 120)
        if header in caps:
            width = min(width, caps[header])
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, width)


def add_review_dropdowns(ws) -> None:
    header_to_col = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        header_to_col[normalize_str(cell.value)] = col_idx

    decision_formula = '"' + ",".join(CLINICIAN_DECISION_CHOICES) + '"'
    dv_decision = DataValidation(
        type="list",
        formula1=decision_formula,
        allow_blank=True,
    )
    ws.add_data_validation(dv_decision)

    for row_idx in range(2, ws.max_row + 1):
        if "clinician_decision" in header_to_col:
            dv_decision.add(ws.cell(row=row_idx, column=header_to_col["clinician_decision"]))


def append_kv_sheet(wb: Workbook, title: str, rows: List[Dict[str, str]], content_cap: int = 110) -> None:
    ws = wb.create_sheet(title)
    ws.append(["section", "content"])
    for row in rows:
        ws.append([row["section"], row["content"]])
    style_header(ws)
    ws.freeze_panes = "A2"
    auto_width_with_caps(ws, caps={"section": 42, "content": content_cap})


# =============================================================================
# README / reference sheet builders
# =============================================================================

def build_readme_sheet_rows() -> List[Dict[str, str]]:
    return [
        {
            "section": "purpose",
            "content": (
                "This workbook is designed for clinicians to validate the outputs of the LLM judges "
                "on two questions: (1) whether a clinical trial is relevant to a patient, and "
                "(2) whether the patient is eligible for that trial. Since a trial may include "
                "multiple subcohorts, the evaluation is conducted in two stages. First, one LLM "
                "judge identifies all subcohorts within the trial that are relevant to the patient. "
                "Second, another LLM judge evaluates the patient’s eligibility for each relevant "
                "subcohort on an individual basis."
            ),
        },
        {
            "section": "task",
            "content": (
                "We ask you to review the LLM judges’ decisions for both relevance and eligibility "
                "and indicate whether you accept or do not accept each decision. We also ask you "
                "to provide a brief explanation for your decision."
            ),
        },
        {
            "section": "guideline",
            "content": (
                "We have left two columns for you to complete: clinician_decision and "
                "clinician_rationale. In clinician_decision, please indicate whether you accept "
                "or do not accept the LLM judge’s decision. In clinician_rationale, please briefly "
                "explain your reasoning.\n\n"
                "Please note that the definition of relevance differs across the three evaluation "
                "objectives. In addition, the eligibility guidelines were designed to make SMT-based "
                "matching as consistent as possible with the LLM judge’s internal assumptions. For "
                "this reason, some of the eligibility criteria may appear to be written from a more "
                "engineering-oriented perspective than a clinician-oriented one."
            ),
        },
        {
            "section": "Goal",
            "content": (
                "The original datasets were often ambiguous and inconsistent in their task "
                "definitions and in the assumptions underlying both relevance and eligibility. "
                "We initially attempted to follow the criteria implicitly assumed in those datasets, "
                "but found that this was not sufficiently reliable. As a result, we evaluate our "
                "system under three settings, ranging from the strictest to the most relaxed, and "
                "show that our approach performs better across all of them. We do not claim that "
                "these settings are necessarily the most clinically representative. Rather, our goal "
                "is to demonstrate that the system performs robustly under multiple reasonable "
                "settings and can be configured to suit different needs."
            ),
        },
        {
            "section": "Why the eligibility judge is more engineering flavored than clinician oriented",
            "content": (
                "Eligibility assessment is often ambiguous, both in how LLMs interpret criteria and "
                "in how implicit assumptions influence their decisions. To make eligibility judgments "
                "more consistent across our system and the baseline, we had to explicitly control for "
                "these assumptions and abstract away or factor out many details. As a result, some of "
                "the eligibility criteria may differ from how clinicians would naturally express or "
                "interpret them in practice."
            ),
        },
        {
            "section": "how_to_read_each_row",
            "content": (
                "Each row should be read from left to right: first the patient and trial context, "
                "then the relevance and eligibility outputs from the LLM judges, and finally the "
                "two columns to be completed by the clinician."
            ),
        },
    ]


def build_column_guide_rows() -> List[Dict[str, str]]:
    return [
        {
            "section": "column_overview",
            "content": (
                "Each review sheet contains one sampled patient-trial pair per row. "
                "The columns provide the patient context, the trial content, the LLM judges’ "
                "outputs, and space for clinician validation."
            ),
        },
        {"section": "sample_id", "content": "Unique identifier for the sampled review row."},
        {
            "section": "retrieval_objective",
            "content": "Clinician-facing name of the retrieval objective under which this patient-trial pair was sampled.",
        },
        {
            "section": "retrieval_objective_description",
            "content": "A short description of the retrieval objective for this review sheet.",
        },
        {"section": "patient_id", "content": "Identifier of the patient case."},
        {"section": "trial_id", "content": "Identifier of the clinical trial."},
        {
            "section": "union_bucket_used_for_sampling",
            "content": "The sampling bucket assigned to this patient-trial pair in the union sampling step.",
        },
        {
            "section": "relevance_definition",
            "content": "The formal definition of relevance used for this retrieval objective.",
        },
        {
            "section": "relevance_instructions",
            "content": "Additional instructions given to the LLM judge for identifying relevant subcohorts.",
        },
        {
            "section": "patient_note",
            "content": "The patient vignette or prescreen note used for evaluation.",
        },
        {
            "section": "trial_listing",
            "content": (
                "A structured presentation of the full trial information for the sampled trial, "
                "including title, summary, criteria, and metadata when available."
            ),
        },
        {
            "section": "any_trial_relevant_from_llm_judge",
            "content": (
                "Whether the relevance LLM judge determined that the trial contains at least one "
                "subcohort relevant to the patient."
            ),
        },
        {
            "section": "relevant_subcohorts",
            "content": "The relevant subcohort(s) identified by the relevance LLM judge, shown in structured form.",
        },
        {
            "section": "any_trial_relevant_and_eligible_from_llm_judge",
            "content": (
                "Whether the eligibility LLM judge determined that the patient is eligible for at least "
                "one subcohort that was identified as relevant."
            ),
        },
        {
            "section": "relevant_subcohorts_eligibility_determination",
            "content": "The eligibility LLM judge’s structured eligibility decisions for the relevant subcohort(s).",
        },
        {
            "section": "relevance_rationale",
            "content": "A short rationale summarizing why the relevance LLM judge considered the trial relevant.",
        },
        {
            "section": "clinician_decision",
            "content": (
                "Column for clinician completion. Enter one of: accept, reject_relevance, "
                "reject_eligibility, reject_both, or uncertain."
            ),
        },
        {
            "section": "clinician_rationale",
            "content": (
                "Column for clinician completion. Briefly explain the reasoning for your decision."
            ),
        },
    ]


def build_decision_guide_rows() -> List[Dict[str, str]]:
    return [
        {
            "section": "fixed_decision_scheme",
            "content": (
                "Please use the following fixed scheme in clinician_decision: "
                "accept, reject_relevance, reject_eligibility, reject_both, or uncertain."
            ),
        },
        {
            "section": "accept",
            "content": "Use accept if you agree with both the relevance and eligibility judgments.",
        },
        {
            "section": "reject_relevance",
            "content": "Use reject_relevance if you disagree with the relevance judgment.",
        },
        {
            "section": "reject_eligibility",
            "content": "Use reject_eligibility if you agree that the trial is relevant but disagree with the eligibility judgment.",
        },
        {
            "section": "reject_both",
            "content": "Use reject_both if you disagree with both the relevance and eligibility judgments.",
        },
        {
            "section": "uncertain",
            "content": "Use uncertain if the case cannot be judged confidently from the provided information.",
        },
        {
            "section": "clinician_rationale_guidance",
            "content": (
                "In clinician_rationale, briefly explain the reason for your decision. "
                "When applicable, specify whether the disagreement is about relevance, eligibility, "
                "or both, and identify the key clinical fact or trial criterion that drove your decision."
            ),
        },
    ]


def build_retrieval_objective_rows() -> List[Dict[str, str]]:
    return [
        {
            "section": "retrieval_objectives_overview",
            "content": (
                "The workbook contains three retrieval objectives that differ in how broadly a trial "
                "is considered a match to the patient case."
            ),
        },
        {
            "section": "1. chief_complaint_treating",
            "content": (
                "Chief-complaint-treating: Trials intended to directly treat the patient's chief complaint. "
                "This is the narrowest retrieval objective."
            ),
        },
        {
            "section": "2. any_condition_treating",
            "content": (
                "Any-condition-treating: Trials intended to treat any clinically important condition "
                "described in the patient case, not only the chief complaint."
            ),
        },
        {
            "section": "3. any_condition_relevant",
            "content": (
                "Any-condition-relevant: Trials that are clinically relevant to any condition "
                "described in the patient case, including non-treatment relationships when appropriate. "
                "This is the broadest retrieval objective."
            ),
        },
        {
            "section": "sheet_mapping",
            "content": (
                "The three review sheets correspond to the three retrieval objectives: "
                "Chief-complaint-treating, Any-condition-treating, and Any-condition-relevant."
            ),
        },
    ]


def build_prompt_rows_for_mode(mode: str, pc: Dict[str, str]) -> List[Dict[str, str]]:
    return [
        {
            "section": f"{MODE_TO_DISPLAY[mode]}: description",
            "content": MODE_TO_DESCRIPTION[mode],
        },
        {
            "section": f"{MODE_TO_DISPLAY[mode]}: relevance_base_prompt",
            "content": pc["relevance_base"] or "N/A",
        },
        {
            "section": f"{MODE_TO_DISPLAY[mode]}: relevance_definition_prompt",
            "content": pc["relevance_def"] or "N/A",
        },
        {
            "section": f"{MODE_TO_DISPLAY[mode]}: relevance_instructions_prompt",
            "content": pc["relevance_instructions"] or "N/A",
        },
        {
            "section": f"{MODE_TO_DISPLAY[mode]}: eligibility_prompt_template",
            "content": pc["eligibility_template"] or "N/A",
        },
    ]


# =============================================================================
# Main
# =============================================================================

def main() -> None:
    ap = argparse.ArgumentParser()

    ap.add_argument(
        "--sample-csv",
        default="./physician_validation_union_sample/sampled_items.csv",
        help="sampled_items.csv produced by sample_physician_validation_union.py",
    )
    ap.add_argument(
        "--output-xlsx",
        default="./physician_validation_union_sample/clinician_review.xlsx",
    )

    ap.add_argument("--patient-corpus", default="../../dataset/clinical_trial/sigir/queries.jsonl")
    ap.add_argument("--trial-corpus", default="../../dataset/clinical_trial/sigir/corpus.jsonl")

    ap.add_argument(
        "--shared-pair-cache-root",
        default="./_shared_pair_cache/pair_cache/.cache",
        help="Root of shared relevance/eligibility pair cache.",
    )

    ap.add_argument("--prompt-dir", default="./prompts")
    args = ap.parse_args()

    sample_csv = Path(args.sample_csv).resolve()
    output_xlsx = Path(args.output_xlsx).resolve()
    patient_corpus = load_jsonl_as_dict(Path(args.patient_corpus).resolve(), id_key="_id")
    trial_corpus = load_jsonl_as_dict(Path(args.trial_corpus).resolve(), id_key="_id")
    shared_pair_cache_root = Path(args.shared_pair_cache_root).resolve()
    prompt_dir = Path(args.prompt_dir).resolve()

    sample_rows = read_csv_rows(sample_csv)
    review_items = build_review_items(sample_rows)
    prompt_components = load_prompt_components(prompt_dir)

    rows_by_mode: Dict[str, List[dict]] = {m: [] for m in MODES}

    for item in review_items:
        if item.mode not in MODES:
            continue

        pc = prompt_components[item.mode]
        patient_note = patient_corpus.get(item.patient_id, {}).get("text", "")
        trial_obj = trial_corpus.get(item.trial_id, {})
        trial_listing = build_trial_listing(trial_obj)

        pair_dir = get_pair_cache_trial_dir(
            shared_pair_cache_root=shared_pair_cache_root,
            mode=item.mode,
            patient_id=item.patient_id,
            trial_id=item.trial_id,
        )

        relevance_path = pair_dir / "relevance.txt"
        eligibility_path = pair_dir / "eligibility.txt"

        relevance_output = read_text(relevance_path) if relevance_path.exists() else ""
        eligibility_output = read_text(eligibility_path) if eligibility_path.exists() else ""

        parsed_rel = parse_relevance_bool(relevance_output)
        parsed_elig = parse_eligibility_bool(eligibility_output)

        relevant_obj = parse_relevant_subcohorts_obj(relevance_output)
        eligibility_obj = parse_eligibility_obj(eligibility_output)
        relevance_rationale = extract_optional_relevance_rationale(
            relevant_subcohorts_obj=relevant_obj,
            relevance_output_raw=relevance_output,
        )

        if parsed_rel is True:
            any_trial_relevant = "yes"
        elif parsed_rel is False:
            any_trial_relevant = "no"
        else:
            any_trial_relevant = "N/A"

        if parsed_rel is True and parsed_elig is True:
            any_trial_relevant_and_eligible = "yes"
        elif parsed_rel is False or (parsed_rel is True and parsed_elig is False):
            any_trial_relevant_and_eligible = "no"
        else:
            any_trial_relevant_and_eligible = "N/A"

        rows_by_mode[item.mode].append({
            "sample_id": item.sample_id or "N/A",
            "retrieval_objective": MODE_TO_DISPLAY.get(item.mode, "N/A"),
            "retrieval_objective_description": MODE_TO_DESCRIPTION.get(item.mode, "N/A"),
            "patient_id": item.patient_id or "N/A",
            "trial_id": item.trial_id or "N/A",
            "union_bucket_used_for_sampling": item.union_bucket_used_for_sampling or "N/A",
            "relevance_definition": pc["relevance_def"] or "N/A",
            "relevance_instructions": pc["relevance_instructions"] or "N/A",
            "patient_note": patient_note.strip() if patient_note.strip() else "N/A",
            "trial_listing": trial_listing if trial_listing.strip() else "N/A",
            "any_trial_relevant_from_llm_judge": any_trial_relevant,
            "relevant_subcohorts": (
                json.dumps(relevant_obj, ensure_ascii=False, indent=2)
                if relevant_obj not in (None, "", [])
                else "N/A"
            ),
            "any_trial_relevant_and_eligible_from_llm_judge": any_trial_relevant_and_eligible,
            "relevant_subcohorts_eligibility_determination": (
                json.dumps(eligibility_obj, ensure_ascii=False, indent=2)
                if eligibility_obj not in (None, "", [])
                else "N/A"
            ),
            "relevance_rationale": relevance_rationale if relevance_rationale.strip() else "N/A",
            "clinician_decision": "",
            "clinician_rationale": "",
        })

    wb = Workbook()
    wb.remove(wb.active)

    review_headers = [
        "sample_id",
        "retrieval_objective",
        "retrieval_objective_description",
        "patient_id",
        "trial_id",
        "union_bucket_used_for_sampling",
        "relevance_definition",
        "relevance_instructions",
        "patient_note",
        "trial_listing",
        "any_trial_relevant_from_llm_judge",
        "relevant_subcohorts",
        "any_trial_relevant_and_eligible_from_llm_judge",
        "relevant_subcohorts_eligibility_determination",
        "relevance_rationale",
        "clinician_decision",
        "clinician_rationale",
    ]

    for mode in MODES:
        ws = wb.create_sheet(MODE_TO_SHEET[mode])
        ws.append(review_headers)
        for row in rows_by_mode.get(mode, []):
            ws.append([row.get(h, "") for h in review_headers])

        style_header(ws)
        style_semantic_columns(ws)
        ws.freeze_panes = "A2"
        add_review_dropdowns(ws)
        auto_width_with_caps(
            ws,
            caps={
                "retrieval_objective": 32,
                "retrieval_objective_description": 52,
                "relevance_definition": 70,
                "relevance_instructions": 70,
                "patient_note": 70,
                "trial_listing": 90,
                "relevant_subcohorts": 60,
                "relevant_subcohorts_eligibility_determination": 60,
                "relevance_rationale": 60,
                "clinician_rationale": 60,
            },
        )

    append_kv_sheet(wb, "README", build_readme_sheet_rows(), content_cap=110)
    append_kv_sheet(wb, "column_guide", build_column_guide_rows(), content_cap=95)
    append_kv_sheet(wb, "decision_guide", build_decision_guide_rows(), content_cap=95)
    append_kv_sheet(wb, "retrieval_objectives", build_retrieval_objective_rows(), content_cap=100)

    append_kv_sheet(
        wb,
        "prompts_ccr",
        build_prompt_rows_for_mode("ccr", prompt_components["ccr"]),
        content_cap=130,
    )
    append_kv_sheet(
        wb,
        "prompts_all",
        build_prompt_rows_for_mode("all", prompt_components["all"]),
        content_cap=130,
    )
    append_kv_sheet(
        wb,
        "prompts_all_explore",
        build_prompt_rows_for_mode("all-explore", prompt_components["all-explore"]),
        content_cap=130,
    )

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)

    print(f"[DONE] wrote workbook: {output_xlsx}")
    for mode in MODES:
        print(f"[INFO] {MODE_TO_SHEET[mode]} rows: {len(rows_by_mode.get(mode, []))}")


if __name__ == "__main__":
    main()